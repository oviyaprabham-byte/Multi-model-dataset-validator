import io
import os
from datetime import datetime
from typing import Dict, Any

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
    HRFlowable, KeepTogether
)
from reportlab.lib.enums import TA_CENTER, TA_LEFT

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from models.schemas import ValidationResult


# ── Color palette ────────────────────────────────────────────────────────
DARK_BG = colors.HexColor("#0F172A")
ACCENT = colors.HexColor("#6366F1")
SUCCESS = colors.HexColor("#10B981")
WARNING = colors.HexColor("#F59E0B")
DANGER = colors.HexColor("#EF4444")
LIGHT_GRAY = colors.HexColor("#F8FAFC")
MID_GRAY = colors.HexColor("#E2E8F0")
TEXT_DARK = colors.HexColor("#1E293B")


def _severity_color(sev: str) -> colors.HexColor:
    if sev == "high":
        return DANGER
    elif sev == "medium":
        return WARNING
    return SUCCESS


def _status_color(status: str) -> colors.HexColor:
    if status == "pass":
        return SUCCESS
    elif status == "fail":
        return DANGER
    return WARNING


def generate_pdf(result: ValidationResult, output_path: str) -> str:
    doc = SimpleDocTemplate(
        output_path,
        pagesize=A4,
        rightMargin=2 * cm,
        leftMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("Title", fontSize=24, textColor=ACCENT,
                                  alignment=TA_CENTER, spaceAfter=6, fontName="Helvetica-Bold")
    subtitle_style = ParagraphStyle("Subtitle", fontSize=11, textColor=TEXT_DARK,
                                     alignment=TA_CENTER, spaceAfter=2)
    h2_style = ParagraphStyle("H2", fontSize=14, textColor=ACCENT,
                               fontName="Helvetica-Bold", spaceBefore=12, spaceAfter=6)
    h3_style = ParagraphStyle("H3", fontSize=11, textColor=TEXT_DARK,
                               fontName="Helvetica-Bold", spaceBefore=8, spaceAfter=4)
    normal_style = ParagraphStyle("Normal", fontSize=9, textColor=TEXT_DARK, spaceAfter=3)

    story = []

    # ── Cover ────────────────────────────────────────────────────────────
    story.append(Spacer(1, 1 * cm))
    story.append(Paragraph("📊 Dataset Validation Report", title_style))
    story.append(Paragraph(f"Generated: {result.timestamp}", subtitle_style))
    story.append(Paragraph(f"Session ID: {result.session_id}", subtitle_style))
    story.append(HRFlowable(width="100%", thickness=2, color=ACCENT, spaceAfter=10))

    # ── Dataset Info ─────────────────────────────────────────────────────
    story.append(Paragraph("Dataset Information", h2_style))
    info = result.dataset_info
    info_data = [
        ["Field", "Value"],
        ["Filename", info.filename],
        ["File Size", info.size_human],
        ["Dataset Type", info.dataset_type.upper()],
        ["Format", info.format or "N/A"],
    ]
    if info.total_rows:
        info_data.append(["Total Rows", str(info.total_rows)])
    if info.total_columns:
        info_data.append(["Total Columns", str(info.total_columns)])
    if info.total_files:
        info_data.append(["Total Files", str(info.total_files)])

    info_table = Table(info_data, colWidths=[6 * cm, 11 * cm])
    info_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), ACCENT),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("BACKGROUND", (0, 1), (-1, -1), LIGHT_GRAY),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT_GRAY]),
        ("GRID", (0, 0), (-1, -1), 0.5, MID_GRAY),
        ("PADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(info_table)
    story.append(Spacer(1, 0.5 * cm))

    # ── Quality Score ─────────────────────────────────────────────────────
    story.append(Paragraph("Quality Score", h2_style))
    qs = result.quality_score
    badge_color = SUCCESS if qs.badge in ("Excellent", "Good") else (WARNING if qs.badge == "Average" else DANGER)
    score_data = [
        [f"Overall Score: {qs.score}/100", f"Rating: {qs.badge}",
         f"Status: {'✓ READY' if result.is_ready else '✗ NOT READY'}"],
    ]
    score_table = Table(score_data, colWidths=[6 * cm, 5 * cm, 6 * cm])
    score_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, 0), ACCENT),
        ("BACKGROUND", (1, 0), (1, 0), badge_color),
        ("BACKGROUND", (2, 0), (2, 0), SUCCESS if result.is_ready else DANGER),
        ("TEXTCOLOR", (0, 0), (-1, -1), colors.white),
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 11),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("PADDING", (0, 0), (-1, -1), 12),
        ("ROUNDEDCORNERS", [4, 4, 4, 4]),
    ]))
    story.append(score_table)
    story.append(Spacer(1, 0.5 * cm))

    # Score breakdown
    breakdown_data = [["Dimension", "Score"]]
    for dim, val in qs.breakdown.items():
        max_map = {
            "missing_values": 20, "duplicates": 15, "corrupted": 25,
            "empty": 15, "format": 10, "class_imbalance": 10, "missing_labels": 5
        }
        max_v = max_map.get(dim, 10)
        breakdown_data.append([dim.replace("_", " ").title(), f"{val}/{max_v}"])
    bd_table = Table(breakdown_data, colWidths=[10 * cm, 7 * cm])
    bd_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), ACCENT),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT_GRAY]),
        ("GRID", (0, 0), (-1, -1), 0.5, MID_GRAY),
        ("PADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(bd_table)
    story.append(Spacer(1, 0.5 * cm))

    # ── Validation Checks ─────────────────────────────────────────────────
    story.append(Paragraph("Validation Checks", h2_style))
    checks_data = [["Check", "Status", "Message"]]
    for c in result.checks:
        status_label = "✓ Pass" if c.status == "pass" else ("✗ Fail" if c.status == "fail" else "⚠ Warn")
        checks_data.append([c.name, status_label, c.message[:80]])

    checks_table = Table(checks_data, colWidths=[5 * cm, 2.5 * cm, 9.5 * cm])
    style_cmds = [
        ("BACKGROUND", (0, 0), (-1, 0), ACCENT),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, MID_GRAY),
        ("PADDING", (0, 0), (-1, -1), 5),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT_GRAY]),
    ]
    for i, c in enumerate(result.checks, start=1):
        col = _status_color(c.status)
        style_cmds.append(("TEXTCOLOR", (1, i), (1, i), col))
        style_cmds.append(("FONTNAME", (1, i), (1, i), "Helvetica-Bold"))
    checks_table.setStyle(TableStyle(style_cmds))
    story.append(checks_table)
    story.append(Spacer(1, 0.5 * cm))

    # ── Issues ────────────────────────────────────────────────────────────
    if result.issues:
        story.append(Paragraph("Issues Found", h2_style))
        issues_data = [["Issue", "Severity", "Description"]]
        for iss in result.issues:
            issues_data.append([iss.title, iss.severity.upper(),
                                 (iss.description or "")[:70]])
        iss_table = Table(issues_data, colWidths=[5 * cm, 2.5 * cm, 9.5 * cm])
        iss_cmds = [
            ("BACKGROUND", (0, 0), (-1, 0), ACCENT),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, MID_GRAY),
            ("PADDING", (0, 0), (-1, -1), 5),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT_GRAY]),
        ]
        for i, iss in enumerate(result.issues, start=1):
            col = _severity_color(iss.severity)
            iss_cmds.append(("TEXTCOLOR", (1, i), (1, i), col))
            iss_cmds.append(("FONTNAME", (1, i), (1, i), "Helvetica-Bold"))
        iss_table.setStyle(TableStyle(iss_cmds))
        story.append(iss_table)
        story.append(Spacer(1, 0.5 * cm))

    # ── Recommendations ───────────────────────────────────────────────────
    if result.recommendations:
        story.append(Paragraph("Recommendations", h2_style))
        for i, rec in enumerate(result.recommendations, 1):
            story.append(Paragraph(f"{i}. {rec.title}", h3_style))
            story.append(Paragraph(rec.description, normal_style))

    # ── Final Status ──────────────────────────────────────────────────────
    story.append(Spacer(1, 0.5 * cm))
    story.append(HRFlowable(width="100%", thickness=2, color=ACCENT))
    story.append(Spacer(1, 0.3 * cm))
    verdict_text = "✓ READY FOR TRAINING" if result.is_ready else "✗ NOT READY FOR TRAINING"
    verdict_style = ParagraphStyle("Verdict", fontSize=18, textColor=SUCCESS if result.is_ready else DANGER,
                                    fontName="Helvetica-Bold", alignment=TA_CENTER, spaceAfter=6)
    story.append(Paragraph(verdict_text, verdict_style))

    if not result.is_ready and result.readiness_reasons:
        story.append(Paragraph("Reasons:", h3_style))
        for r in result.readiness_reasons:
            story.append(Paragraph(f"• {r}", normal_style))

    doc.build(story)
    return output_path


def generate_excel(result: ValidationResult, output_path: str) -> str:
    wb = openpyxl.Workbook()

    header_fill = PatternFill("solid", fgColor="6366F1")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    pass_fill = PatternFill("solid", fgColor="D1FAE5")
    fail_fill = PatternFill("solid", fgColor="FEE2E2")
    warn_fill = PatternFill("solid", fgColor="FEF3C7")
    high_fill = PatternFill("solid", fgColor="FEE2E2")
    med_fill = PatternFill("solid", fgColor="FEF3C7")
    low_fill = PatternFill("solid", fgColor="D1FAE5")
    border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    def style_header(ws, row=1):
        for cell in ws[row]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

    def style_cell(cell, fill=None):
        if fill:
            cell.fill = fill
        cell.border = border
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # ── Sheet 1: Summary ─────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Summary"
    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 40
    ws.append(["Field", "Value"])
    style_header(ws)

    info = result.dataset_info
    rows = [
        ("Report Date", result.timestamp),
        ("Session ID", result.session_id),
        ("Filename", info.filename),
        ("File Size", info.size_human),
        ("Dataset Type", info.dataset_type.upper()),
        ("Quality Score", f"{result.quality_score.score}/100"),
        ("Rating", result.quality_score.badge),
        ("Training Status", "READY" if result.is_ready else "NOT READY"),
    ]
    if info.total_rows:
        rows.append(("Total Rows", info.total_rows))
    if info.total_columns:
        rows.append(("Total Columns", info.total_columns))
    if info.total_files:
        rows.append(("Total Files", info.total_files))

    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            style_cell(cell)

    # ── Sheet 2: Validation Checks ────────────────────────────────────────
    ws2 = wb.create_sheet("Validation Checks")
    ws2.column_dimensions["A"].width = 25
    ws2.column_dimensions["B"].width = 12
    ws2.column_dimensions["C"].width = 50
    ws2.column_dimensions["D"].width = 12
    ws2.append(["Check Name", "Status", "Message", "Severity"])
    style_header(ws2)

    for i, c in enumerate(result.checks, start=2):
        ws2.append([c.name, c.status.upper(), c.message, c.severity or ""])
        fill = pass_fill if c.status == "pass" else (fail_fill if c.status == "fail" else warn_fill)
        for cell in ws2[i]:
            style_cell(cell, fill)

    # ── Sheet 3: Issues ────────────────────────────────────────────────────
    ws3 = wb.create_sheet("Issues")
    ws3.column_dimensions["A"].width = 25
    ws3.column_dimensions["B"].width = 12
    ws3.column_dimensions["C"].width = 55
    ws3.column_dimensions["D"].width = 10
    ws3.append(["Issue Title", "Severity", "Description", "Count"])
    style_header(ws3)

    for i, iss in enumerate(result.issues, start=2):
        ws3.append([iss.title, iss.severity.upper(), iss.description, iss.count or ""])
        fill = high_fill if iss.severity == "high" else (med_fill if iss.severity == "medium" else low_fill)
        for cell in ws3[i]:
            style_cell(cell, fill)

    # ── Sheet 4: Recommendations ──────────────────────────────────────────
    ws4 = wb.create_sheet("Recommendations")
    ws4.column_dimensions["A"].width = 5
    ws4.column_dimensions["B"].width = 30
    ws4.column_dimensions["C"].width = 60
    ws4.column_dimensions["D"].width = 12
    ws4.append(["#", "Title", "Description", "Priority"])
    style_header(ws4)

    for i, rec in enumerate(result.recommendations, start=1):
        ws4.append([i, rec.title, rec.description, rec.priority.upper()])
        fill = high_fill if rec.priority == "high" else (med_fill if rec.priority == "medium" else low_fill)
        for cell in ws4[i + 1]:
            style_cell(cell, fill)

    # ── Sheet 5: Score Breakdown ──────────────────────────────────────────
    ws5 = wb.create_sheet("Score Breakdown")
    ws5.column_dimensions["A"].width = 25
    ws5.column_dimensions["B"].width = 15
    ws5.column_dimensions["C"].width = 15
    ws5.append(["Dimension", "Score Earned", "Max Score"])
    style_header(ws5)

    max_map = {
        "missing_values": 20, "duplicates": 15, "corrupted": 25,
        "empty": 15, "format": 10, "class_imbalance": 10, "missing_labels": 5
    }
    for i, (dim, val) in enumerate(result.quality_score.breakdown.items(), start=2):
        max_v = max_map.get(dim, 10)
        ws5.append([dim.replace("_", " ").title(), val, max_v])
        fill = pass_fill if val >= max_v * 0.7 else (warn_fill if val >= max_v * 0.4 else fail_fill)
        for cell in ws5[i]:
            style_cell(cell, fill)

    wb.save(output_path)
    return output_path
